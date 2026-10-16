import os
import openpyxl
from reports import formatar_relatorio_texto, formatar_notas_texto, gerar_planilha_excel

def test_formatar_relatorio_texto_completo():
    """Valida a formatação de relatório com SAs e anotações do dia."""
    registros = [
        {"sa_id": "96569", "data": "2026-10-16", "horas": 3.0, "descricao": "Refatoração de código"},
        {"sa_id": "96569", "data": "2026-10-16", "horas": 1.5, "descricao": "Testes unitários"},
        {"sa_id": "99888", "data": "2026-10-16", "horas": 2.0, "descricao": "Ajuste na API"}
    ]
    notas = [
        {"hora": "09:00", "conteudo": "Daily meeting com time"},
        {"hora": "14:30", "conteudo": "Ajudou colega com debug"}
    ]

    texto = formatar_relatorio_texto(registros, titulo="Rastro Diário", notas=notas)
    assert "Total de Horas nas SAs:* `6.50h`" in texto
    assert "SA: 96569* (`4.50h`)" in texto
    assert "SA: 99888* (`2.00h`)" in texto
    assert "Daily meeting com time" in texto
    assert "Ajudou colega com debug" in texto

def test_formatar_relatorio_texto_vazio():
    """Valida relatório quando não há apontamentos nem notas."""
    texto = formatar_relatorio_texto([], titulo="Rastro Diário", notas=[])
    assert "Nenhum registro ou anotação para este período." in texto

def test_formatar_notas_texto():
    """Valida formatação exclusiva de bloco de notas."""
    notas = [{"hora": "11:00", "conteudo": "Anotação rápida"}]
    texto = formatar_notas_texto(notas, "16/10/2026")
    assert "Anotação rápida" in texto
    assert "[11:00]" in texto

    texto_vazio = formatar_notas_texto([], "16/10/2026")
    assert "Nenhuma anotação registrada hoje." in texto_vazio

def test_gerar_planilha_excel(tmp_path):
    """Valida a criação e estrutura da planilha Excel gerada."""
    caminho_xlsx = str(tmp_path / "teste_horas.xlsx")
    registros = [
        {"data": "2026-10-16", "sa_id": "96569", "horas": 4.0, "descricao": "Implementação do bot"},
        {"data": "2026-10-16", "sa_id": "99888", "horas": 3.5, "descricao": "Reunião de alinhamento"}
    ]

    gerar_planilha_excel(registros, caminho_xlsx)
    assert os.path.exists(caminho_xlsx)

    wb = openpyxl.load_workbook(caminho_xlsx)
    ws = wb.active
    assert ws.title == "Lançamento de Horas"

    # Verificar cabeçalhos
    assert ws.cell(row=1, column=1).value == "Data"
    assert ws.cell(row=1, column=2).value == "SA / Demanda"
    assert ws.cell(row=1, column=3).value == "Horas Gastas"
    assert ws.cell(row=1, column=4).value == "Descrição da Atividade"

    # Verificar primeira linha de dados
    assert ws.cell(row=2, column=1).value == "2026-10-16"
    assert ws.cell(row=2, column=2).value == "96569"
    assert ws.cell(row=2, column=3).value == 4.0

    # Verificar totalizador
    linha_total = len(registros) + 2
    assert ws.cell(row=linha_total, column=2).value == "TOTAL:"
    assert ws.cell(row=linha_total, column=3).value == 7.5

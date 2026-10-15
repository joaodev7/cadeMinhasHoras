from typing import List, Dict, Any, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def formatar_relatorio_texto(
    registros: List[Dict[str, Any]], 
    titulo: str = "Relatório de Horas",
    notas: Optional[List[Dict[str, Any]]] = None
) -> str:
    """Gera um relatório formatado em Markdown com SAs e as anotações do bloco de notas."""
    if not registros and not notas:
        return f"📅 *{titulo}*\nNenhum registro ou anotação para este período."

    linhas = [f"📊 *{titulo}*\n"]

    # Seção 1: Horas nas SAs
    if registros:
        por_sa: Dict[str, List[Dict[str, Any]]] = {}
        total_geral = 0.0

        for reg in registros:
            sa = reg["sa_id"]
            if sa not in por_sa:
                por_sa[sa] = []
            por_sa[sa].append(reg)
            total_geral += reg["horas"]

        linhas.append(f"⏱️ *Total de Horas nas SAs:* `{total_geral:.2f}h`\n")

        for sa, itens in por_sa.items():
            total_sa = sum(i["horas"] for i in itens)
            linhas.append(f"📌 *SA: {sa}* (`{total_sa:.2f}h`)")
            for i in itens:
                linhas.append(f"  • *{i['data']}* (`{i['horas']}h`): {i['descricao']}")
            linhas.append("")
    else:
        linhas.append("⏱️ *Horas nas SAs:* Nenhuma hora apontada em SA ainda.\n")

    # Seção 2: Bloco de Notas (atividades sem SA)
    if notas:
        linhas.append("📝 *Bloco de Notas do Dia (Sem SA):*")
        for n in notas:
            linhas.append(f"  • `[{n['hora']}]` {n['conteudo']}")
        linhas.append("")

    return "\n".join(linhas)

def formatar_notas_texto(notas: List[Dict[str, Any]], data_formatada: str) -> str:
    """Formata apenas o bloco de notas para o comando /notas."""
    if not notas:
        return f"📝 *Bloco de Notas ({data_formatada})*\nNenhuma anotação registrada hoje."

    linhas = [f"📝 *Bloco de Notas ({data_formatada})*:\n"]
    for n in notas:
        linhas.append(f"• `[{n['hora']}]` {n['conteudo']}")
    return "\n".join(linhas)

def gerar_planilha_excel(registros: List[Dict[str, Any]], caminho_saida: str) -> str:
    """Gera uma planilha Excel estilizada pronta para envio ou cópia das SAs."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Lançamento de Horas"

    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    headers = ["Data", "SA / Demanda", "Horas Gastas", "Descrição da Atividade"]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    total_horas = 0.0
    for row_idx, reg in enumerate(registros, start=2):
        ws.append([
            reg["data"],
            reg["sa_id"],
            reg["horas"],
            reg["descricao"]
        ])
        total_horas += reg["horas"]

        ws.cell(row=row_idx, column=1).alignment = center_align
        ws.cell(row=row_idx, column=2).alignment = center_align
        ws.cell(row=row_idx, column=3).alignment = center_align
        ws.cell(row=row_idx, column=4).alignment = left_align

        for col_num in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col_num).font = data_font
            ws.cell(row=row_idx, column=col_num).border = thin_border

    last_row = len(registros) + 2
    ws.cell(row=last_row, column=2, value="TOTAL:")
    ws.cell(row=last_row, column=2).font = Font(bold=True)
    ws.cell(row=last_row, column=2).alignment = Alignment(horizontal="right")

    total_cell = ws.cell(row=last_row, column=3, value=total_horas)
    total_cell.font = Font(bold=True)
    total_cell.alignment = center_align

    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 50

    wb.save(caminho_saida)
    return caminho_saida
